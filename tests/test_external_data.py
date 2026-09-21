import csv
import json
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

from participacion.adapters.tabular.csv_reader import read_csv_table
from participacion.adapters.tabular.xlsx_reader import read_xlsx_table
from participacion.application.external_data.mapping import (
    FieldMapping,
    ExternalDataMapping,
    MappingStatus,
    map_table,
)
from participacion.application.external_data.models import AvailabilityStatus
from participacion.application.external_data.validation import ValidationCode


class ExternalDataTests(unittest.TestCase):
    def mapping(self, **metrics):
        return ExternalDataMapping(
            participant=FieldMapping("participant", "string"),
            session=FieldMapping("session", "string"),
            metrics=metrics,
        )

    def test_csv_reader_returns_neutral_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.csv"
            path.write_text("participant,session,minutes\nuser-1,s-1,12\n", encoding="utf-8")
            table = read_csv_table(path)

        self.assertEqual(table.columns, ("participant", "session", "minutes"))
        self.assertEqual(table.rows[0].row_number, 2)
        self.assertEqual(table.rows[0].values["minutes"], "12")
        self.assertEqual(table.source_name, "source.csv")

    def test_xlsx_reader_selects_requested_sheet_and_preserves_basic_types(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "attendance"
            sheet.append(["participant", "session", "minutes", "day"])
            sheet.append(["user-1", "s-1", 12, date(2026, 1, 5)])
            workbook.create_sheet("other")
            workbook.save(path)
            workbook.close()

            table = read_xlsx_table(path, sheet_name="attendance")

        self.assertEqual(table.rows[0].values["minutes"], 12)
        self.assertEqual(table.rows[0].values["day"], datetime(2026, 1, 5))
        self.assertEqual(table.sheet_name, "attendance")

    def test_xlsx_reader_rejects_missing_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            workbook.create_sheet("attendance")
            workbook.save(path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "Hoja inexistente"):
                read_xlsx_table(path, sheet_name="missing")

    def test_mapping_creates_canonical_facts_with_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.csv"
            path.write_text("participant,session,minutes\nuser-1,s-1,12\n", encoding="utf-8")
            table = read_csv_table(path)
            result = map_table(table, self.mapping(minutes=FieldMapping("minutes", "integer")))

        self.assertEqual(result.status, MappingStatus.VALID)
        self.assertEqual(len(result.facts), 1)
        fact = result.facts[0]
        self.assertEqual((fact.participant_external_id, fact.session_external_id), ("user-1", "s-1"))
        self.assertEqual((fact.metric, fact.value), ("minutes", 12))
        self.assertEqual(fact.availability, AvailabilityStatus.AVAILABLE)
        self.assertEqual((fact.provenance.file_name, fact.provenance.sheet_name, fact.provenance.row_number),
                         ("source.csv", None, 2))

    def test_missing_required_column_is_invalid(self):
        table = read_csv_table_from_rows(["participant", "session"], [["user-1", "s-1"]])
        result = map_table(table, self.mapping(minutes=FieldMapping("missing", "integer", required=True)))
        self.assertEqual(result.status, MappingStatus.INVALID)
        self.assertTrue(any(issue.code is ValidationCode.MISSING_COLUMN for issue in result.issues))

    def test_invalid_type_is_needs_review_and_marks_fact_invalid(self):
        table = read_csv_table_from_rows(["participant", "session", "minutes"], [["user-1", "s-1", "bad"]])
        result = map_table(table, self.mapping(minutes=FieldMapping("minutes", "integer")))
        self.assertEqual(result.status, MappingStatus.NEEDS_REVIEW)
        self.assertEqual(result.facts[0].availability, AvailabilityStatus.INVALID)
        self.assertTrue(any(issue.code is ValidationCode.INVALID_VALUE for issue in result.issues))

    def test_incomplete_row_is_needs_review_without_inventing_identity(self):
        table = read_csv_table_from_rows(["participant", "session", "minutes"], [["", "s-1", "12"]])
        result = map_table(table, self.mapping(minutes=FieldMapping("minutes", "integer")))
        self.assertEqual(result.status, MappingStatus.NEEDS_REVIEW)
        self.assertEqual(result.facts, ())
        self.assertTrue(any(issue.code is ValidationCode.INCOMPLETE_ROW for issue in result.issues))

    def test_duplicate_fact_key_is_needs_review_and_not_overwritten(self):
        table = read_csv_table_from_rows(
            ["participant", "session", "minutes"],
            [["user-1", "s-1", "12"], ["user-1", "s-1", "13"]],
        )
        result = map_table(table, self.mapping(minutes=FieldMapping("minutes", "integer")))
        self.assertEqual(result.status, MappingStatus.NEEDS_REVIEW)
        self.assertEqual(len(result.facts), 2)
        self.assertTrue(any(issue.code is ValidationCode.DUPLICATE_FACT for issue in result.issues))

    def test_ambiguous_mapping_is_invalid(self):
        table = read_csv_table_from_rows(
            ["participant", "session", "value"], [["user-1", "s-1", "1"]]
        )
        mapping = ExternalDataMapping(
            participant=FieldMapping("value", "string"),
            session=FieldMapping("session", "string"),
            metrics={"metric_a": FieldMapping("value", "string")},
        )
        result = map_table(table, mapping)
        self.assertEqual(result.status, MappingStatus.INVALID)
        self.assertTrue(any(issue.code is ValidationCode.AMBIGUOUS_MAPPING for issue in result.issues))

    def test_example_configuration_is_data_not_adapter_logic(self):
        config = json.loads(
            Path("configs/pilot.example.json").read_text(encoding="utf-8")
        )
        self.assertEqual(config["attendance"]["source"]["type"], "google_sheets")
        parsed = ExternalDataMapping.from_dict(config["attendance"]["mapping"])
        self.assertEqual(parsed.participant.column, "external_participant_id")
        self.assertEqual(parsed.session.column, "class_date")
        table = read_csv_table_from_rows(
            ["external_participant_id", "class_date", "attendance_value_a", "attendance_value_b"],
            [["external-user", "2026-01-01", "0.5", "0.6"]],
        )
        result = map_table(table, parsed)
        self.assertEqual(result.status, MappingStatus.VALID)
        self.assertEqual({fact.metric for fact in result.facts}, {
            "attendance_metric_a", "attendance_metric_b",
        })

    def test_generic_adapter_has_no_basf_names(self):
        for path in Path("src/participacion/adapters/tabular").glob("*.py"):
            text = path.read_text(encoding="utf-8")
            self.assertNotIn("BASF", text)
            self.assertNotIn("gOS_", text)
            self.assertNotIn("BAJAS", text)

    def test_core_has_no_external_ingestion_dependencies(self):
        for path in Path("src/participacion/core").glob("*.py"):
            text = path.read_text(encoding="utf-8")
            for forbidden in ("external_data", "tabular", "openpyxl", "import csv"):
                self.assertNotIn(forbidden, text, path)


def read_csv_table_from_rows(headers, rows):
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "synthetic.csv"
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)
        return read_csv_table(path)


if __name__ == "__main__":
    unittest.main()
