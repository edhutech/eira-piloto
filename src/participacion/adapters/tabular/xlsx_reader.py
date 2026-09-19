from __future__ import annotations

from pathlib import Path

from ...application.external_data.models import TabularRow, TabularTable
from .csv_reader import TabularReadError


def read_xlsx_table(path: str | Path, *, sheet_name: str) -> TabularTable:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("XLSX requiere openpyxl instalado; usa participacion-agent[xlsx]") from exc
    source = Path(path).expanduser()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise TabularReadError(f"Hoja inexistente: {sheet_name}")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise TabularReadError(f"La hoja está vacía: {sheet_name}")
        columns = _headers(header)
        values = tuple(
            TabularRow(number, dict(zip(columns, row)))
            for number, row in enumerate(rows, 2)
            if any(value is not None and str(value).strip() for value in row)
        )
        return TabularTable(source.name, sheet_name, columns, values)
    finally:
        workbook.close()


def _headers(row: tuple[object, ...]) -> tuple[str, ...]:
    columns = tuple(str(value).strip() if value is not None and str(value).strip() else f"__unnamed_{index}"
                    for index, value in enumerate(row, 1))
    if len(set(columns)) != len(columns):
        raise TabularReadError("La hoja contiene encabezados duplicados")
    return columns
